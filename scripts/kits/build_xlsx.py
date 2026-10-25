"""Construit les classeurs Excel des kits Import-eXPert.
Sources : Guide des procédures douanières Maroc, éd. 2026 (à jour au 04/10/2026),
et « Import Process Management » (maj 17/10/2026)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

OUT = '/home/claude/kits/out'
os.makedirs(OUT, exist_ok=True)

NAVY, RED, GOLD = '153E82', 'A82D3C', 'C9A84C'
F = 'Arial'
thin = Side(style='thin', color='C9D2E3')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill('solid', fgColor=NAVY)
IN_FILL = PatternFill('solid', fgColor='FFF7D6')     # cellules à saisir
CALC_FILL = PatternFill('solid', fgColor='F2F5FA')   # cellules calculées
EX_FILL = PatternFill('solid', fgColor='EAF6EE')     # ligne d'exemple
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
DISCLAIMER = ("Outil d'aide à la conformité. Il ne se substitue ni aux textes officiels ni à une consultation. "
              "Données à jour au 04/10/2026 (Guide des procédures douanières Maroc, éd. 2026) : vérifier avant toute opération "
              "que le texte cité n'a pas été modifié.")


def title(ws, text, sub, width_cols):
    ws['A1'] = 'IMPORT-eXPERT'
    ws['A1'].font = Font(name=F, bold=True, color=RED, size=10)
    ws['A2'] = text
    ws['A2'].font = Font(name=F, bold=True, color=NAVY, size=16)
    ws['A3'] = sub
    ws['A3'].font = Font(name=F, italic=True, color='5A6475', size=9)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=width_cols)
    ws['A3'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[3].height = 30


def header(ws, row, cols, widths):
    for i, (h, w) in enumerate(zip(cols, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=F, bold=True, color='FFFFFF', size=10)
        c.fill = HDR_FILL
        c.alignment = CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 42
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def style_row(ws, row, ncols, input_cols, fill_override=None, height=None):
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.border = BORDER
        c.alignment = WRAP
        is_in = col in input_cols
        c.font = Font(name=F, size=10, color='0000FF' if is_in else '000000')
        c.fill = fill_override or (IN_FILL if is_in else CALC_FILL)
    if height:
        ws.row_dimensions[row].height = height


def legend(ws, row, extra=()):
    lines = [('Cellules jaunes, texte bleu', 'à saisir'), ('Cellules grises', 'calculées : ne pas modifier'),
             ('Ligne verte', "exemple fictif, à remplacer ou supprimer")] + list(extra)
    ws.cell(row=row, column=1, value='Légende').font = Font(name=F, bold=True, color=NAVY)
    for i, (a, b) in enumerate(lines, start=1):
        ws.cell(row=row + i, column=1, value=a).font = Font(name=F, size=9, bold=True)
        ws.cell(row=row + i, column=2, value=b).font = Font(name=F, size=9)


def guide_sheet(wb, name, titre, paragraphs):
    ws = wb.create_sheet(name, 0)
    title(ws, titre, DISCLAIMER, 6)
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 110
    r = 5
    for kind, text in paragraphs:
        c = ws.cell(row=r, column=2, value=text)
        if kind == 'h':
            c.font = Font(name=F, bold=True, color=NAVY, size=12)
            r += 1
        else:
            c.font = Font(name=F, size=10)
            c.alignment = Alignment(wrap_text=True, vertical='top')
            ws.row_dimensions[r].height = max(15, 15 * (len(text) // 120 + 1))
            r += 1
    ws.sheet_view.showGridLines = False
    return ws


# ═══════════════════════════════ KIT 01 — Matrice d'éligibilité ═══════════════════════════════
def kit01():
    wb = Workbook()
    wb.remove(wb.active)

    # Référentiel des accords
    acc = wb.create_sheet('Accords')
    title(acc, 'Référentiel des accords et des règles d’origine',
          'Une ligne par accord. Les cellules « À vérifier » signalent un point que le Guide ne fixe pas : consulter la circulaire citée avant usage.', 12)
    cols = ['Accord', 'Pays couverts', 'Règles d’origine applicables', 'Preuves admises', 'Déclaration d’origine sans agrément',
            'Validité de la preuve (mois)', 'Tolérance / critère de valeur', 'Cumul', 'No-drawback', 'Mentions à l’export (Maroc)',
            'Conservation des justificatifs (années)', 'Sources']
    header(acc, 5, cols, [16, 22, 34, 28, 22, 12, 30, 34, 26, 30, 14, 30])
    rows = [
        ['UE', 'États membres de l’Union européenne', 'Convention pan-euro-méditerranéenne révisée, application provisoire du 01/01/2026 au 31/12/2027 au plus tard',
         'EUR.1 ou déclaration d’origine sur facture / document commercial (EUR-MED supprimés)', 'Envoi ≤ 6 000 EUR ; sans limite pour l’exportateur agréé',
         10, '15 % du poids net (ch. 2 à 24) ou du prix départ usine (autres), hors textiles', 'Étendu à toutes les parties appliquant les règles ; bilatéral seulement pour les ch. 50 à 63',
         'Limité aux ch. 50 à 63 ; inapplicable aux échanges bilatéraux', '« REVISED RULES » en case 7 ; « CUMULATION APPLIED WITH … » en cas de cumul', 3,
         'Circ. n° 6706/233 du 31/12/2025, annexe I, art. 5, 7, 17-31 ; circ. n° 6722/233'],
        ['Turquie', 'Türkiye', 'Règles révisées, application provisoire depuis le 01/01/2026 (jusqu’au 31/12/2027 au plus tard)',
         'EUR.1 ou déclaration d’origine (EUR-MED supprimés)', 'Envoi ≤ 6 000 EUR ; sans limite pour l’exportateur agréé', 10,
         '15 % (comme UE)', 'Diagonal Maroc–UE–Turquie maintenu sous règles identiques', 'Comme UE', '« REVISED RULES » ; « CUMULATION APPLIED WITH … »', 3,
         'Circ. n° 6706/233 ; circ. n° 6722/233'],
        ['AELE', 'Suisse, Norvège, Islande, Liechtenstein', 'Anciennes règles maintenues jusqu’à la modification du protocole',
         'À vérifier (circ. n° 4980/233)', 'À vérifier (circ. n° 4980/233)', None, 'À vérifier (circ. n° 4980/233)', 'À vérifier',
         'À vérifier', 'Selon circ. n° 4980/233', None,
         'Circ. n° 4980/233 citée par circ. n° 6706/233 ; produits agricoles transformés : circ. n° 6409/222 (seul l’élément agricole est taxé)'],
        ['Agadir', 'Tunisie, Égypte, Jordanie', 'Anciennes règles maintenues jusqu’à la modification des protocoles',
         'À vérifier (circ. n° 5047/223)', 'À vérifier (circ. n° 5047/223)', None, 'À vérifier (circ. n° 5047/223)', 'À vérifier', 'À vérifier',
         'Selon circ. n° 5047/223', None, 'Circ. n° 5047/223 citée par circ. n° 6706/233'],
        ['États-Unis', 'États-Unis d’Amérique', 'ALE en vigueur depuis le 01/01/2006 ; règles spécifiques de l’annexe IX ou, à défaut, article nouveau + critère de 35 %',
         'Aucun certificat : l’importateur qui demande la préférence est réputé certifier l’éligibilité', 'Sans objet', None,
         '(VMP + CDT) / VEP ≥ 35 % ; de minimis textile 7 % du poids du composant déterminant', 'Total entre le Maroc et les États-Unis',
         'Sans objet', 'Transport direct entre les parties', 5, 'Loi n° 28-04 ; circ. n° 4977/222, section IV ; circ. n° 6705/222'],
        ['Royaume-Uni', 'Royaume-Uni', 'Accord d’association appliqué depuis le 01/01/2021 (reprise de l’accord Maroc–UE) ; décision n° 01/2026 élargissant le cumul',
         'EUR.1 ou déclaration sur facture (cumul bilatéral) ; EUR-MED (cumul diagonal)', 'À vérifier (incidence des règles révisées UE)', None, 'À vérifier',
         'Bilatéral : UE, Norvège, Islande ; diagonal : Suisse, Turquie, Égypte, Tunisie, Jordanie ; cumul élargi par la décision n° 01/2026',
         'Non en cumul bilatéral ; oui en cumul diagonal', 'Code BADR « UK »', None, 'Circ. n° 6139/222 ; circ. n° 6758/233 du 07/07/2026'],
        ['ZLECAf', 'Pays africains appliquant l’accord', 'Règles de l’accord et de sa circulaire d’application : à vérifier pour chaque produit',
         'Certificat d’origine visé par l’autorité du pays d’exportation ; délivrance dématérialisée sur BADR à l’export', 'Sans objet', None, 'À vérifier',
         'À vérifier', 'À vérifier', 'Certificat ZLECAf (BADR)', None, 'Note D87/2026/ADII/AC/100 ; circ. n° 6530/223 ; brochure BADR'],
        ['Droit commun', 'Pays sans accord préférentiel', 'Origine non préférentielle : entièrement obtenu ou transformation complète',
         'Certificat de droit commun « CO Rose » à l’export', 'Sans objet', None,
         'Plus-value ≥ valeur d’importation des produits mis en œuvre, ou opération de l’annexe I du décret', 'Sans objet', 'Sans objet', 'Sans objet', None,
         'CDII, art. 16 à 19 ; décret n° 2-77-862, art. 1er et 2 ; circ. n° 5549/230'],
    ]
    for i, r in enumerate(rows, start=6):
        for j, v in enumerate(r, start=1):
            acc.cell(row=i, column=j, value=v)
        style_row(acc, i, len(cols), input_cols=set(), height=78)
        for j in range(1, len(cols) + 1):
            acc.cell(row=i, column=j).fill = PatternFill('solid', fgColor='FFFFFF')
        acc.cell(row=i, column=1).font = Font(name=F, bold=True, color=NAVY)
    last_acc = 5 + len(rows)

    # Analyse produit
    an = wb.create_sheet('Analyse produit')
    title(an, 'Analyse d’éligibilité produit par produit',
          'Saisir le produit, l’accord du pays de destination et les valeurs. Le plafond de la règle de liste se lit dans la liste de l’accord '
          '(annexe II de la circ. n° 6706/233 pour l’UE et la Turquie). Pour les États-Unis, renseigner VMP + CDT et VEP.', 17)
    cols = ['Réf. produit', 'Désignation', 'Code SH (10 chiffres)', 'Accord', 'Prix départ usine (DH)', 'Valeur des matières non originaires (DH)',
            'Plafond de la règle de liste (% max. non originaires)', '% de matières non originaires', 'Règle de liste (critère de valeur)',
            'Tolérance à examiner', 'États-Unis : VMP + CDT (DH)', 'États-Unis : VEP (DH)', 'États-Unis : ratio', 'États-Unis : critère de 35 %',
            'Droit commun : plus-value', 'Preuve à produire', 'Verdict']
    header(an, 5, cols, [11, 26, 15, 13, 14, 15, 15, 12, 20, 24, 13, 13, 10, 18, 22, 32, 30])
    acc_list = f"Accords!$A$6:$A${last_acc}"
    dv = DataValidation(type='list', formula1=f'={acc_list}', allow_blank=True)
    an.add_data_validation(dv)
    first, last = 6, 45
    for r in range(first, last + 1):
        an[f'H{r}'] = f'=IF(OR(E{r}="",E{r}=0,F{r}=""),"",F{r}/E{r})'
        an[f'I{r}'] = (f'=IF(H{r}="","",IF(G{r}="","Plafond non renseigné",IF(H{r}<=G{r},"Conforme","Non conforme")))')
        an[f'J{r}'] = (f'=IF(AND(I{r}="Non conforme",OR(D{r}="UE",D{r}="Turquie")),'
                       f'"Examiner la tolérance de 15 % (hors textiles)","")')
        an[f'M{r}'] = f'=IF(OR(L{r}="",L{r}=0,K{r}=""),"",K{r}/L{r})'
        an[f'N{r}'] = f'=IF(D{r}<>"États-Unis","—",IF(M{r}="","Saisir VMP+CDT et VEP",IF(M{r}>=0.35,"≥ 35 % : conforme","< 35 % : non conforme")))'
        an[f'O{r}'] = (f'=IF(D{r}<>"Droit commun","—",IF(OR(E{r}="",F{r}=""),"Saisir les valeurs",'
                       f'IF(E{r}-F{r}>=F{r},"Plus-value suffisante","Insuffisante : vérifier l’annexe I du décret")))')
        an[f'P{r}'] = f'=IF(D{r}="","",IFERROR(INDEX(Accords!$D$6:$D${last_acc},MATCH(D{r},{acc_list},0)),""))'
        an[f'Q{r}'] = (f'=IF(D{r}="","",IF(D{r}="États-Unis",N{r},IF(D{r}="Droit commun",O{r},'
                       f'IF(I{r}="Conforme","Éligible sous réserve de la règle de liste complète",'
                       f'IF(I{r}="Non conforme","Non éligible en l’état",'
                       f'"Renseigner le plafond de la règle de liste")))))')
        an[f'H{r}'].number_format = '0.0%'
        an[f'G{r}'].number_format = '0%'
        an[f'M{r}'].number_format = '0.0%'
        for c in 'EFKL':
            an[f'{c}{r}'].number_format = '#,##0'
        dv.add(f'D{r}')
        style_row(an, r, len(cols), input_cols={1, 2, 3, 4, 5, 6, 7, 11, 12})
    # exemple fictif
    ex = ['EX-01', 'Faisceaux de câbles (exemple fictif)', '8544.30.00.00', 'UE', 1_000_000, 380_000, 0.40]
    for j, v in enumerate(ex, start=1):
        an.cell(row=first, column=j, value=v)
    style_row(an, first, len(cols), input_cols={1, 2, 3, 4, 5, 6, 7, 11, 12}, fill_override=EX_FILL)
    an.cell(row=first, column=7).comment = Comment(
        'Valeur fictive pour l’exemple : lire le plafond réel dans la règle de liste de la position du produit fini '
        '(annexe II, circ. n° 6706/233).', 'Import-eXPert')
    an.conditional_formatting.add(f'I{first}:I{last}', CellIsRule(operator='equal', formula=['"Non conforme"'], font=Font(color=RED, bold=True)))
    an.conditional_formatting.add(f'I{first}:I{last}', CellIsRule(operator='equal', formula=['"Conforme"'], font=Font(color='1E7346', bold=True)))
    legend(an, last + 2, [('Règle de liste', 'le calcul ne couvre que le critère de valeur ; les règles de changement de position ou d’opération spécifique s’apprécient à part')])

    # Suivi des preuves
    sp = wb.create_sheet('Suivi des preuves')
    title(sp, 'Suivi des preuves d’origine reçues et émises',
          'Validité : 10 mois pour l’UE et la Turquie (annexe I, art. 23) ; déclaration à long terme du fournisseur : 2 ans au plus (art. 29). '
          'Conservation : 3 ans au moins (UE, Turquie), 5 ans pour les États-Unis.', 12)
    cols = ['Réf. opération', 'Sens (Import / Export)', 'Partenaire', 'Accord', 'Type de preuve', 'N° de la preuve', 'Date d’établissement',
            'Validité (mois)', 'Valable jusqu’au', 'Jours restants', 'Statut', 'Conserver jusqu’au']
    header(sp, 5, cols, [13, 12, 22, 13, 26, 16, 14, 11, 14, 11, 20, 15])
    dv_acc = DataValidation(type='list', formula1=f'={acc_list}', allow_blank=True)
    dv_sens = DataValidation(type='list', formula1='"Import,Export"', allow_blank=True)
    dv_type = DataValidation(type='list', formula1='"EUR.1,Déclaration d’origine,Déclaration fournisseur long terme,Certificat ZLECAf,Certification importateur (États-Unis),Certificat CO Rose"', allow_blank=True)
    for d in (dv_acc, dv_sens, dv_type):
        sp.add_data_validation(d)
    for r in range(6, 106):
        sp[f'H{r}'] = (f'=IF(E{r}="","",IF(E{r}="Déclaration fournisseur long terme",24,'
                       f'IFERROR(IF(INDEX(Accords!$F$6:$F${last_acc},MATCH(D{r},{acc_list},0))="","",'
                       f'INDEX(Accords!$F$6:$F${last_acc},MATCH(D{r},{acc_list},0))),"")))')
        sp[f'I{r}'] = f'=IF(OR(G{r}="",H{r}=""),"",EDATE(G{r},H{r}))'
        sp[f'J{r}'] = f'=IF(I{r}="","",I{r}-TODAY())'
        sp[f'K{r}'] = (f'=IF(G{r}="","",IF(H{r}="","Validité : à vérifier",IF(J{r}<0,"EXPIRÉE",'
                       f'IF(J{r}<=30,"Expire sous 30 jours","Valide"))))')
        sp[f'L{r}'] = (f'=IF(G{r}="","",IFERROR(EDATE(G{r},12*INDEX(Accords!$K$6:$K${last_acc},MATCH(D{r},{acc_list},0))),"À vérifier"))')
        for c in 'GIL':
            sp[f'{c}{r}'].number_format = 'DD/MM/YYYY'
        dv_acc.add(f'D{r}'); dv_sens.add(f'B{r}'); dv_type.add(f'E{r}')
        style_row(sp, r, len(cols), input_cols={1, 2, 3, 4, 5, 6, 7})
    import datetime as dt
    for j, v in enumerate(['OP-2026-001', 'Import', 'Fournisseur fictif (Espagne)', 'UE', 'EUR.1', 'A 0000000', dt.date(2026, 6, 15)], start=1):
        sp.cell(row=6, column=j, value=v)
    style_row(sp, 6, len(cols), input_cols={1, 2, 3, 4, 5, 6, 7}, fill_override=EX_FILL)
    sp['G6'].number_format = 'DD/MM/YYYY'
    sp.conditional_formatting.add('K6:K105', CellIsRule(operator='equal', formula=['"EXPIRÉE"'], font=Font(color=RED, bold=True), fill=PatternFill('solid', fgColor='FCEEF0')))
    sp.conditional_formatting.add('K6:K105', CellIsRule(operator='equal', formula=['"Expire sous 30 jours"'], font=Font(color='9A620C', bold=True)))

    guide_sheet(wb, 'Mode d’emploi', 'Kit 01 · Matrice d’éligibilité à l’origine préférentielle', [
        ('h', 'À quoi sert ce classeur'),
        ('p', "Vérifier, produit par produit et accord par accord, si une marchandise peut bénéficier d'un droit préférentiel, quelle preuve produire, et suivre la validité et la conservation de ces preuves."),
        ('h', 'Mode d’emploi'),
        ('p', "1. Onglet « Accords » : référentiel des règles par accord. Les mentions « À vérifier » signalent un point non tranché par le Guide : consultez la circulaire citée."),
        ('p', "2. Onglet « Analyse produit » : choisissez l'accord du pays de destination (ou d'origine, à l'import), saisissez le prix départ usine et la valeur des matières non originaires, puis le plafond de la règle de liste applicable. Pour les États-Unis, saisissez VMP + CDT et VEP : le critère de 35 % est calculé."),
        ('p', "3. Onglet « Suivi des preuves » : enregistrez chaque EUR.1, déclaration d'origine ou déclaration de fournisseur ; la date limite de validité et la date de fin de conservation sont calculées."),
        ('h', 'Les règles à retenir'),
        ('p', "Il faut toujours appliquer les règles de l'accord invoqué, à la date de l'opération : un même produit peut être originaire pour l'UE et ne pas l'être pour les États-Unis (Guide, n° 2-021)."),
        ('p', "Depuis le 01/01/2026, UE et Turquie suivent les règles révisées (tolérance de 15 %, preuves valables 10 mois, EUR-MED supprimés, mention « REVISED RULES ») ; AELE et Agadir conservent les anciennes règles (circ. n° 6706/233)."),
        ('p', "Une déclaration d'origine peut être établie par tout exportateur jusqu'à 6 000 EUR, et sans limite par un exportateur agréé (annexe I, art. 18 et 19)."),
        ('p', "Le calcul de la feuille « Analyse produit » ne couvre que le critère de valeur. Une règle de liste peut aussi exiger un changement de position tarifaire ou une opération spécifique : ce point s'apprécie à partir de la liste de l'accord."),
        ('h', 'Sources'),
        ('p', "Guide des procédures douanières Maroc, éd. 2026, n° 2-020 à 2-028. Circ. ADII n° 6706/233, 6722/233, 6758/233, 4977/222, 6139/222, 6409/222, 6530/223, 5549/230 ; notes D87/2026, D99/2026, D105/2026."),
    ])
    wb.active = 0
    wb.save(f'{OUT}/Kit01_Matrice_eligibilite_origine.xlsx')


# ═══════════════════════════════ KIT 02 — Tableau d'apurement ═══════════════════════════════
def kit02():
    wb = Workbook()
    wb.remove(wb.active)
    par = wb.create_sheet('Paramètres')
    title(par, 'Délais de séjour par régime', 'Délais réglementaires (Guide, n° 7-013, 7-021, 2-004, 7-016). Le délai initial peut être prorogé dans la limite de la durée maximale.', 5)
    header(par, 5, ['Régime', 'Délai initial (mois)', 'Durée maximale (mois)', 'Autorité de prorogation', 'Source'], [44, 14, 14, 40, 40])
    regs = [
        ['ATPA', 6, 24, 'Administration ; au-delà de 2 ans : dérogation', 'Décret, art. 105 ; CDII, art. 137'],
        ['ATPA artisans', 3, 3, '—', 'Circ. n° 6164/313, annexe II'],
        ['ATPA retouches', 6, 6, '—', 'Circ. n° 6164/313, annexes I et II'],
        ['ATPA cession (reliquat < 12 mois)', 12, 12, '—', 'Circ. n° 6164/313, annexe I'],
        ['AT matériels Ca, Cc, D3, D9, D11', 24, 48, 'Ministre ou directeur général délégué (jusqu’au double)', 'Décret, art. 132'],
        ['AT matériels Cb', 12, 24, 'Ministre ou directeur général délégué (jusqu’au double)', 'Décret, art. 132'],
        ['AT autres biens catégorie D', 6, 12, 'Ministre ou directeur général délégué (jusqu’au double)', 'Décret, art. 132'],
        ['AT expositions', 6, 6, '—', 'Circ. n° 6348/313, annexe I'],
        ['AT essais', 24, 24, '—', 'Circ. n° 6348/313, annexe I'],
        ['AT échantillons (unités exportatrices)', 6, 6, '—', 'Circ. n° 6348/313, annexe I'],
        ['AT commerce triangulaire hors enceinte', 3, 3, '—', 'Note D885/21/ADII/100'],
        ['Transformation sous douane (envois fractionnés)', 3, 12, 'Prorogation de 9 mois au plus', 'Décret, art. 172 septies'],
        ['Compensation par équivalence / exportation préalable', 24, 24, '—', 'CDII, art. 142-4° et 5°'],
    ]
    for i, r in enumerate(regs, start=6):
        for j, v in enumerate(r, start=1):
            par.cell(row=i, column=j, value=v)
        style_row(par, i, 5, input_cols=set())
        for j in range(1, 6):
            par.cell(row=i, column=j).fill = PatternFill('solid', fgColor='FFFFFF')
    lr = 5 + len(regs)
    r0 = lr + 2
    consts = [('Seuil d’alerte 1 (jours)', 90, 'Alerte J-90 (IPM, annexe D.3)'),
              ('Seuil d’alerte 2 (jours)', 60, 'Alerte J-60 : prorogation et exportation hors délai relèvent d’autorités différentes (Guide, n° 7-006)'),
              ('Plafond de mise à la consommation sans amende', 0.15, '15 % des quantités exportées sous ATPA (circ. n° 6348/313, annexe I)'),
              ('Solde échu exonéré (DH)', 500, 'Soldes échus inférieurs à 500 DH (CDII, art. 116 bis-2°)')]
    par.cell(row=r0, column=1, value='Constantes').font = Font(name=F, bold=True, color=NAVY)
    for k, (a, b, c) in enumerate(consts, start=1):
        par.cell(row=r0 + k, column=1, value=a)
        par.cell(row=r0 + k, column=2, value=b)
        par.cell(row=r0 + k, column=4, value=c)
        style_row(par, r0 + k, 4, input_cols={2})
    par.cell(row=r0 + 3, column=2).number_format = '0%'
    A1, A2, MAC, SOLDE = f'Paramètres!$B${r0+1}', f'Paramètres!$B${r0+2}', f'Paramètres!$B${r0+3}', f'Paramètres!$B${r0+4}'
    REG = f'Paramètres!$A$6:$A${lr}'

    # Comptes
    co = wb.create_sheet('Comptes')
    title(co, 'Comptes sous régime économique : échéances, apurement, alertes',
          'Une ligne par compte (une déclaration de placement). Les imputations se saisissent dans l’onglet « Imputations ».', 21)
    cols = ['N° compte (DUM)', 'Régime', 'Bureau', 'Date d’enregistrement', 'Intrant', 'Unité', 'Quantité placée', 'Droits et taxes suspendus (DH)',
            'Taux de rendement (intrant par unité de produit)', 'Prorogation accordée (mois)', 'Délai initial (mois)', 'Durée maximale (mois)',
            'Date d’échéance', 'Jours restants', 'Quantité d’intrants imputée', 'Solde d’intrants', 'Taux d’apurement', 'Droits sur le solde (DH)',
            'Mise à la consommation / exportations', 'Statut', 'Action recommandée']
    header(co, 5, cols, [16, 26, 14, 14, 20, 8, 12, 14, 14, 12, 10, 10, 13, 10, 14, 12, 11, 14, 14, 18, 44])
    dv = DataValidation(type='list', formula1=f'={REG}', allow_blank=True)
    co.add_data_validation(dv)
    IMP = 'Imputations'
    for r in range(6, 106):
        co[f'K{r}'] = f'=IF(B{r}="","",IFERROR(INDEX(Paramètres!$B$6:$B${lr},MATCH(B{r},{REG},0)),""))'
        co[f'L{r}'] = f'=IF(B{r}="","",IFERROR(INDEX(Paramètres!$C$6:$C${lr},MATCH(B{r},{REG},0)),""))'
        co[f'M{r}'] = f'=IF(OR(D{r}="",K{r}=""),"",EDATE(D{r},MIN(K{r}+N(J{r}),L{r})))'
        co[f'N{r}'] = f'=IF(M{r}="","",M{r}-TODAY())'
        co[f'O{r}'] = f'=IF(A{r}="","",SUMIFS({IMP}!$H$6:$H$505,{IMP}!$B$6:$B$505,A{r}))'
        co[f'P{r}'] = f'=IF(A{r}="","",G{r}-O{r})'
        co[f'Q{r}'] = f'=IF(OR(A{r}="",N(G{r})=0),"",O{r}/G{r})'
        co[f'R{r}'] = f'=IF(OR(A{r}="",N(G{r})=0),"",MAX(0,P{r})/G{r}*H{r})'
        co[f'S{r}'] = (f'=IF(A{r}="","",IFERROR(SUMIFS({IMP}!$E$6:$E$505,{IMP}!$B$6:$B$505,A{r},{IMP}!$D$6:$D$505,"Mise à la consommation")'
                       f'/SUMIFS({IMP}!$E$6:$E$505,{IMP}!$B$6:$B$505,A{r},{IMP}!$D$6:$D$505,"Exportation"),""))')
        co[f'T{r}'] = (f'=IF(A{r}="","",IF(P{r}<=0,"Apuré",IF(M{r}="","Régime à choisir",IF(N{r}<0,"ÉCHU",'
                       f'IF(N{r}<={A2},"Alerte J-60",IF(N{r}<={A1},"Alerte J-90","En cours"))))))')
        co[f'U{r}'] = (f'=IF(A{r}="","",IF(T{r}="Apuré","Demander le certificat de décharge et vérifier sa réception par la banque",'
                       f'IF(T{r}="ÉCHU",IF(R{r}<{SOLDE},"Solde < 500 DH : exonéré (art. 116 bis) ; régulariser la clôture",'
                       f'"Régulariser sans délai : exportation hors délai (directeur provincial) ou mise à la consommation ; compte échu = facilités suspendues"),'
                       f'IF(T{r}="Alerte J-60","Demander la prorogation (dans la durée maximale ; au-delà : Administration centrale) ou accélérer l’apurement",'
                       f'IF(T{r}="Alerte J-90","Planifier l’apurement et régulariser les déchets taxables pendant la validité du compte",'
                       f'IF(AND(S{r}<>"",S{r}>{MAC}),"Mise à la consommation > 15 % des exportations : la dispense d’amende tombe",""))))))')
        co[f'D{r}'].number_format = 'DD/MM/YYYY'
        co[f'M{r}'].number_format = 'DD/MM/YYYY'
        for c in 'GOP':
            co[f'{c}{r}'].number_format = '#,##0.00'
        for c in 'HR':
            co[f'{c}{r}'].number_format = '#,##0'
        co[f'Q{r}'].number_format = '0%'
        co[f'S{r}'].number_format = '0.0%'
        dv.add(f'B{r}')
        style_row(co, r, len(cols), input_cols={1, 2, 3, 4, 5, 6, 7, 8, 9, 10})
    import datetime as dt
    ex = ['EX-ATPA-001', 'ATPA', 'Tanger Med', dt.date(2026, 5, 4), 'Fil de cuivre (exemple fictif)', 'kg', 40000, 1_250_000, 0.8, 0]
    for j, v in enumerate(ex, start=1):
        co.cell(row=6, column=j, value=v)
    style_row(co, 6, len(cols), input_cols={1, 2, 3, 4, 5, 6, 7, 8, 9, 10}, fill_override=EX_FILL)
    co['D6'].number_format = 'DD/MM/YYYY'
    co['I6'].comment = Comment('Exemple : 0,8 kg de cuivre par faisceau exporté. Le taux déclaré est réputé admis si l’administration ne le conteste pas dans les 2 mois suivant la déclaration d’exportation (CDII, art. 136-2°).', 'Import-eXPert')
    co.conditional_formatting.add('T6:T105', CellIsRule(operator='equal', formula=['"ÉCHU"'], font=Font(color='FFFFFF', bold=True), fill=PatternFill('solid', fgColor=RED)))
    co.conditional_formatting.add('T6:T105', CellIsRule(operator='equal', formula=['"Alerte J-60"'], font=Font(color=RED, bold=True), fill=PatternFill('solid', fgColor='FCEEF0')))
    co.conditional_formatting.add('T6:T105', CellIsRule(operator='equal', formula=['"Alerte J-90"'], font=Font(color='9A620C', bold=True), fill=PatternFill('solid', fgColor='FCF4E3')))
    co.conditional_formatting.add('T6:T105', CellIsRule(operator='equal', formula=['"Apuré"'], font=Font(color='1E7346', bold=True)))

    # Imputations
    im = wb.create_sheet('Imputations')
    title(im, 'Fiches d’imputation : sorties du régime', 'Une ligne par déclaration de sortie. Chaque fiche est jointe à la déclaration de sortie (décret, art. 107).', 9)
    cols = ['Date', 'N° compte', 'N° DUM de sortie', 'Destination', 'Quantité de produits compensateurs', 'Quantité d’intrants en l’état',
            'Taux de rendement (repris du compte)', 'Intrants imputés', 'Observations']
    header(im, 5, cols, [12, 16, 18, 24, 14, 14, 14, 14, 36])
    dv_dest = DataValidation(type='list', formula1='"Exportation,Exportation en l’état,Mise en entrepôt,Admission temporaire,Cession,Mise à la consommation,Destruction ou abandon,Déchets régularisés"', allow_blank=True)
    dv_cpt = DataValidation(type='list', formula1='=Comptes!$A$6:$A$105', allow_blank=True)
    im.add_data_validation(dv_dest); im.add_data_validation(dv_cpt)
    for r in range(6, 506):
        im[f'G{r}'] = f'=IF(B{r}="","",IFERROR(INDEX(Comptes!$I$6:$I$105,MATCH(B{r},Comptes!$A$6:$A$105,0)),""))'
        im[f'H{r}'] = f'=IF(B{r}="","",N(E{r})*N(G{r})+N(F{r}))'
        im[f'A{r}'].number_format = 'DD/MM/YYYY'
        for c in 'EFH':
            im[f'{c}{r}'].number_format = '#,##0.00'
        dv_dest.add(f'D{r}'); dv_cpt.add(f'B{r}')
        style_row(im, r, len(cols), input_cols={1, 2, 3, 4, 5, 6, 9})
    exs = [[dt.date(2026, 7, 10), 'EX-ATPA-001', 'DUM export 1 (fictive)', 'Exportation', 20000, 0, None, None, 'Premier envoi'],
           [dt.date(2026, 9, 2), 'EX-ATPA-001', 'DUM export 2 (fictive)', 'Exportation', 15000, 0, None, None, 'Deuxième envoi'],
           [dt.date(2026, 9, 20), 'EX-ATPA-001', 'DUM MAC (fictive)', 'Mise à la consommation', 2000, 0, None, None, 'Vente locale autorisée']]
    for i, row in enumerate(exs, start=6):
        for j, v in enumerate(row, start=1):
            if j in (7, 8):
                continue
            im.cell(row=i, column=j, value=v)
        style_row(im, i, len(cols), input_cols={1, 2, 3, 4, 5, 6, 9}, fill_override=EX_FILL)
        im.cell(row=i, column=1).number_format = 'DD/MM/YYYY'

    guide_sheet(wb, 'Mode d’emploi', 'Kit 02 · Tableau d’apurement des comptes sous régime économique', [
        ('h', 'À quoi sert ce classeur'),
        ('p', "Suivre chaque compte ouvert sous régime économique (ATPA, admission temporaire, transformation sous douane…) : échéance, quantités imputées, solde, droits en jeu et alertes, pour ne jamais laisser un compte échu."),
        ('h', 'Mode d’emploi'),
        ('p', "1. « Paramètres » : délais par régime et constantes (alertes J-90 / J-60, plafond de 15 %, seuil de 500 DH). Modifiez une constante si un texte change."),
        ('p', "2. « Comptes » : saisissez chaque déclaration de placement, le taux de rendement (quantité d'intrant par unité de produit compensateur) et, le cas échéant, la prorogation obtenue. Échéance, solde, droits sur le solde et statut sont calculés."),
        ('p', "3. « Imputations » : saisissez chaque déclaration de sortie. Les produits compensateurs sont convertis en intrants avec le taux de rendement ; les intrants sortis en l'état s'ajoutent directement."),
        ('h', 'Les règles à retenir'),
        ('p', "ATPA : séjour initial de 6 mois, prorogeable par l'administration dans la limite de 2 ans à compter de l'enregistrement (décret, art. 105 ; CDII, art. 137)."),
        ('p', "Le compte s'apure globalement ou par apurements partiels, au moyen de fiches d'imputation ; chaque certificat partiel libère la caution à proportion (décret, art. 107 et 108 ; CDII, art. 117)."),
        ('p', "Les déchets taxables se régularisent pendant la validité du compte, sinon les droits sont liquidés d'office (circ. n° 5458/313). Les soldes échus inférieurs à 500 DH sont exonérés (CDII, art. 116 bis-2°)."),
        ('p', "La mise à la consommation est dispensée d'amende dans la limite de 15 % des quantités exportées sous ATPA (circ. n° 6348/313, annexe I)."),
        ('p', "Défaut d'apurement dans les délais : contravention de 2e classe, amende d'une fois et demie les droits (CDII, art. 294 et 294 bis). Un compte échu suspend les facilités de cautionnement (note n° 12014/313)."),
        ('p', "Exportation hors délai : directeurs provinciaux, préfectoraux, des MEAD et des importations-exportations ; prorogation au-delà du délai réglementaire : Administration centrale (circ. n° 6348/313, annexe II)."),
        ('h', 'Sources'),
        ('p', "Guide des procédures douanières Maroc, éd. 2026, n° 7-001 à 7-025 ; Import Process Management, annexe D.3."),
    ])
    wb.active = 0
    wb.save(f'{OUT}/Kit02_Tableau_apurement_RED.xlsx')


# ═══════════════════════════════ KIT 03 — Registre SH ═══════════════════════════════
def kit03_registre():
    wb = Workbook()
    wb.remove(wb.active)
    rg = wb.create_sheet('Registre')
    title(rg, 'Registre centralisé des codes SH',
          'Une ligne par article. Le registre est la pièce maîtresse en cas de contrôle a posteriori : il montre le raisonnement, la source et la date de validation de chaque classement.', 20)
    cols = ['Réf. article', 'Désignation commerciale', 'Description technique (composition, présentation, usage)', 'Code SH (10 chiffres)', 'Contrôle du format',
            'RGI appliquées', 'Justification du classement', 'Source (circulaire, décision anticipée, avis)', 'Type de décision', 'Date de la décision',
            'Valable jusqu’au', 'Validé par (import)', 'Validé par (technique)', 'Date de validation', 'Prochaine revue', 'Jours avant revue',
            'Fiche technique / analyse au dossier', 'Documents et contrôles requis', 'Statut', 'Observations']
    header(rg, 5, cols, [11, 22, 34, 15, 13, 12, 34, 26, 20, 12, 13, 14, 14, 12, 12, 10, 14, 26, 22, 26])
    dv_type = DataValidation(type='list', formula1='"Décision anticipée - classement (5 ans),Décision de classement (art. 15),Circulaire de classement (série 232),Analyse interne"', allow_blank=True)
    dv_on = DataValidation(type='list', formula1='"Oui,Non"', allow_blank=True)
    rg.add_data_validation(dv_type); rg.add_data_validation(dv_on)
    for r in range(6, 306):
        rg[f'E{r}'] = (f'=IF(D{r}="","",IF(AND(LEN(SUBSTITUTE(D{r},".",""))=10,ISNUMBER(--SUBSTITUTE(D{r},".",""))),"10 chiffres : OK","Format à corriger"))')
        rg[f'K{r}'] = (f'=IF(OR(J{r}="",I{r}=""),"",IF(I{r}="Décision anticipée - classement (5 ans)",EDATE(J{r},60),'
                       f'IF(I{r}="Décision de classement (art. 15)","Jusqu’à décision modificative ou changement de nomenclature",'
                       f'IF(I{r}="Circulaire de classement (série 232)","Tant que la circulaire est en vigueur","Revue annuelle"))))')
        rg[f'O{r}'] = f'=IF(N{r}="","",EDATE(N{r},12))'
        rg[f'P{r}'] = f'=IF(O{r}="","",O{r}-TODAY())'
        rg[f'S{r}'] = (f'=IF(A{r}="","",IF(E{r}="Format à corriger","Code à corriger",IF(OR(L{r}="",M{r}=""),"Double validation manquante",'
                       f'IF(Q{r}<>"Oui","Fiche technique manquante",IF(AND(ISNUMBER(K{r}),K{r}<TODAY()),"Décision expirée",'
                       f'IF(AND(P{r}<>"",P{r}<0),"Revue en retard","Conforme"))))))')
        for c in 'JKNO':
            rg[f'{c}{r}'].number_format = 'DD/MM/YYYY'
        dv_type.add(f'I{r}'); dv_on.add(f'Q{r}')
        style_row(rg, r, len(cols), input_cols={1, 2, 3, 4, 6, 7, 8, 9, 10, 12, 13, 14, 17, 18, 20})
    import datetime as dt
    ex = ['EX-001', 'Diffuseur ultrasonique', 'Appareil à ultrasons diffusant des huiles essentielles ; fonction principale : diffusion (exemple tiré du Guide)',
          '8479.89.70.00', None, 'RGI 1 et 6', 'Classement selon la fonction principale ; humidificateurs et appareils médicaux exclus',
          'Circ. ADII n° 6710/232 du 20/01/2026', 'Circulaire de classement (série 232)', dt.date(2026, 1, 20), None,
          'Resp. import', 'Resp. technique', dt.date(2026, 2, 1), None, None, 'Oui', 'À compléter selon la liste DGC et les contrôles techniques', None, '']
    for j, v in enumerate(ex, start=1):
        if v is not None:
            rg.cell(row=6, column=j, value=v)
    style_row(rg, 6, len(cols), input_cols={1, 2, 3, 4, 6, 7, 8, 9, 10, 12, 13, 14, 17, 18, 20}, fill_override=EX_FILL, height=60)
    for c in 'JN':
        rg[f'{c}6'].number_format = 'DD/MM/YYYY'
    rg.conditional_formatting.add('S6:S305', CellIsRule(operator='equal', formula=['"Conforme"'], font=Font(color='1E7346', bold=True)))
    rg.conditional_formatting.add('S6:S305', FormulaRule(formula=['AND(S6<>"",S6<>"Conforme")'], font=Font(color=RED, bold=True)))

    guide_sheet(wb, 'Mode d’emploi', 'Kit 03 · Registre centralisé des codes SH', [
        ('h', 'À quoi sert ce registre'),
        ('p', "Tenir, pour chaque article importé, le classement retenu, son raisonnement RGI par RGI, la source qui le fonde et sa date de validation. C'est la première pièce demandée lors d'un contrôle a posteriori, qui peut remonter sur quatre ans (CDII, art. 99 bis)."),
        ('h', 'Mode d’emploi'),
        ('p', "1. Une ligne par article. Le code tarifaire marocain compte dix chiffres : les six premiers relèvent du SH, les suivants sont des subdivisions nationales (Guide, n° 2-001). La colonne « Contrôle du format » vérifie la longueur."),
        ('p', "2. Indiquez les RGI appliquées et la justification. Recherchez d'abord une circulaire de classement (série 232) portant sur un produit voisin (Guide, n° 2-002)."),
        ('p', "3. Faites valider chaque classement par un responsable import et un responsable technique (double validation ; Import Process Management, 23.2). Une revue est proposée chaque année."),
        ('p', "4. Pour une décision anticipée de classement, la validité de 5 ans est calculée (CDII, art. 45 ter-2°)."),
        ('h', 'Les règles à retenir'),
        ('p', "Une décision ne vaut que pour un produit identique : comparez votre produit à la description de la circulaire, critère par critère (Guide, n° 2-005)."),
        ('p', "Une décision anticipée rendue sur des indications fausses, inexactes ou incomplètes est nulle rétroactivement (CDII, art. 45 ter-2°)."),
        ('p', "Mettez à jour le registre à chaque circulaire « études tarifaires » : un changement de nomenclature peut rendre une sous-position obsolète sans changer le droit (Guide, n° 2-001)."),
        ('h', 'Sources'),
        ('p', "Guide des procédures douanières Maroc, éd. 2026, n° 2-001 à 2-005, 2-060, 2-061 ; CDII, art. 15, 45 ter, 81-2° ; circulaires ADII série 232."),
    ])
    wb.active = 0
    wb.save(f'{OUT}/Kit03_Registre_codes_SH.xlsx')


# ═══════════════════════════════ KIT 03 — Auto-diagnostic ═══════════════════════════════
QUESTIONS = [
    ('Classement', 'Chaque article importé a une sous-position à 10 chiffres documentée (RGI appliquées, justification).', 3, 'CDII, art. 15 ; Guide n° 2-002', 'Constituer un dossier de classement par article (registre SH du kit 03).'),
    ('Classement', 'Tout nouvel article fait l’objet d’une double validation (import et technique).', 2, 'IPM, 23.2', 'Instaurer la double validation à la création de l’article.'),
    ('Classement', 'Les circulaires de classement (série 232) portant sur des produits voisins sont recherchées avant de classer.', 2, 'Guide n° 2-002, 2-005', 'Intégrer la recherche des circulaires 232 à la procédure de classement.'),
    ('Classement', 'La base articles est mise à jour le jour d’entrée en vigueur de chaque circulaire « études tarifaires ».', 2, 'Guide n° 2-001', 'S’abonner à la veille des circulaires tarifaires et dater chaque mise à jour.'),
    ('Classement', 'Les notes de chapitre et les notes complémentaires nationales sont vérifiées avant de choisir la sous-position.', 2, 'CDII, art. 2 et 5-3° ; Guide n° 2-003', 'Ajouter la lecture des notes à la check-list de classement.'),
    ('Classement', 'Les produits dont le classement dépend de la composition disposent d’une fiche technique ou d’un bulletin d’analyse.', 2, 'CDII, art. 81-2°', 'Exiger la fiche technique du fournisseur ; prévoir une analyse pour les cas sensibles.'),
    ('Classement', 'Les décisions anticipées sont suivies avec leur date d’expiration (5 ans classement, 3 ans origine, 1 an valeur).', 2, 'CDII, art. 45 ter-2°', 'Tenir un registre des décisions anticipées avec alertes.'),
    ('Classement', 'Les matériels livrés en plusieurs envois sont déclarés sous autorisation (envois fractionnés ou déclaration globale).', 1, 'CDII, art. 15-4° et 76 bis-4° ; Guide n° 2-004', 'Demander l’autorisation avant l’arrivée du premier envoi.'),
    ('Origine', 'Une fiche d’origine existe par produit et par accord (intrants, origines, règle de liste, calcul).', 3, 'Guide n° 2-021, 2-024', 'Établir les fiches d’origine (matrice du kit 01).'),
    ('Origine', 'Les règles révisées UE–Turquie (depuis le 01/01/2026) sont appliquées : EUR-MED supprimés, mention « REVISED RULES ».', 2, 'Circ. n° 6706/233', 'Réviser les fiches d’origine UE et Turquie et les modèles de preuves.'),
    ('Origine', 'Les preuves d’origine reçues sont contrôlées avant signature de la DUM (validité, mentions, cachet).', 3, 'Circ. n° 6706/233, annexe I, art. 23 ; note D105/2026', 'Ajouter le contrôle des preuves au rapprochement avant signature.'),
    ('Origine', 'Les déclarations d’origine au-delà de 6 000 EUR ne sont établies que par un exportateur agréé.', 2, 'Annexe I, art. 18 et 19', 'Demander des EUR.1 ou le statut d’exportateur agréé.'),
    ('Origine', 'Les déclarations à long terme des fournisseurs sont reçues et suivies (2 ans au plus).', 1, 'Annexe I, art. 29', 'Exiger une déclaration à long terme dès le premier envoi.'),
    ('Origine', 'Les fournisseurs se sont engagés par écrit à répondre aux demandes de vérification de l’origine.', 2, 'Guide n° 2-027', 'Insérer une clause de coopération dans les contrats fournisseurs.'),
    ('Origine', 'Importations des États-Unis : le dossier d’origine du fournisseur (critère de 35 %) est conservé 5 ans.', 1, 'Circ. n° 4977/222', 'Obtenir et archiver le dossier d’origine avant de demander la préférence.'),
    ('Valeur et déclaration', 'La valeur est reconstituée selon l’Incoterm (fret, assurance, redevances, apports) avant signature.', 3, 'IPM, annexe D.2', 'Formaliser une fiche de calcul de la valeur par opération.'),
    ('Valeur et déclaration', 'La facture, la DUM, le titre d’importation et la preuve d’origine sont rapprochés avant signature.', 3, 'IPM, 23.2', 'Mettre en place le contrôle de rapprochement à quatre pièces.'),
    ('Valeur et déclaration', 'Une revue post-mainlevée (J+7 à J+30) permet de révéler les erreurs dans les 30 jours.', 3, 'CDII, art. 78-3°', 'Instaurer la revue post-mainlevée et la procédure de révélation volontaire.'),
    ('Régimes économiques', 'Un échéancier des comptes avec alertes à J-90 et J-60 est tenu.', 3, 'Guide n° 7-006', 'Utiliser le tableau d’apurement du kit 02.'),
    ('Régimes économiques', 'Le sommier BADR est rapproché chaque mois de la comptabilité matières.', 2, 'CDII, art. 116 ter', 'Planifier le rapprochement mensuel.'),
    ('Régimes économiques', 'Les taux de rendement et de déchets sont documentés par produit compensateur.', 2, 'CDII, art. 136', 'Documenter les nomenclatures de fabrication et les taux de déchets.'),
    ('Régimes économiques', 'Les déchets taxables sont régularisés pendant la validité du compte.', 2, 'Circ. n° 5458/313', 'Traiter les déchets au fil de l’eau.'),
    ('Régimes économiques', 'Le ratio de mise à la consommation est suivi sous 15 % des quantités exportées.', 1, 'Circ. n° 6348/313, annexe I', 'Suivre le ratio en continu.'),
    ('Régimes économiques', 'Aucun compte échu non régularisé n’est en cours.', 3, 'Note n° 12014/313 ; CDII, art. 294', 'Régulariser sans délai : un compte échu suspend les facilités de cautionnement.'),
    ('Archivage et communication', 'Chaque opération dispose d’un dossier unique archivé 4 ans (facture, transport, titre, paiement).', 3, 'CDII, art. 42-2° et 72', 'Organiser l’archivage par déclaration.'),
    ('Archivage et communication', 'Les originaux restitués par le transitaire le sont contre décharge.', 1, 'Circ. n° 6639/312', 'Faire signer une décharge à chaque restitution.'),
    ('Archivage et communication', 'Une procédure de réponse écrite aux demandes de communication existe.', 2, 'CDII, art. 297 et 301', 'Désigner un responsable et un délai interne de réponse.'),
    ('Gouvernance', 'Un responsable de la fonction douane est désigné, avec une matrice des responsabilités (RACI).', 2, 'IPM, 23.1', 'Désigner le propriétaire du processus et formaliser le RACI.'),
    ('Gouvernance', 'Les mandats des transitaires sont à jour sur PortNet et visent les acquits à caution si nécessaire.', 2, 'Circ. n° 6421/411 et 5496/400', 'Revoir les mandats et leurs pouvoirs spécifiques.'),
    ('Gouvernance', 'Un audit interne annuel porte sur un échantillon de dossiers.', 2, 'IPM, 23.2', 'Programmer l’audit annuel.'),
    ('Gouvernance', 'Une veille par séries de circulaires et un plan de formation par fonction existent.', 1, 'IPM, 23.1', 'Mettre en place la veille et le plan de formation.'),
    ('Gouvernance', 'Aucune transaction ni condamnation douanière dans les 3 dernières années.', 2, 'CDII, art. 257 ter', 'En cas d’antécédent, renforcer les contrôles : la récidive double le maximum de l’amende.'),
]


def kit03_diag():
    wb = Workbook()
    wb.remove(wb.active)
    q = wb.create_sheet('Questionnaire')
    title(q, 'Auto-diagnostic : exposition au contrôle a posteriori',
          'Répondez Oui, Partiel, Non ou Sans objet. Le score pondère chaque question (1 à 3). Un « Sans objet » est exclu du calcul.', 9)
    cols = ['N°', 'Domaine', 'Question', 'Poids', 'Réponse', 'Points', 'Points max.', 'Action recommandée', 'Référence']
    header(q, 5, cols, [5, 18, 60, 7, 11, 8, 9, 46, 26])
    dv = DataValidation(type='list', formula1='"Oui,Partiel,Non,Sans objet"', allow_blank=True)
    q.add_data_validation(dv)
    for i, (dom, txt, w, ref, act) in enumerate(QUESTIONS, start=1):
        r = 5 + i
        q.cell(row=r, column=1, value=i)
        q.cell(row=r, column=2, value=dom)
        q.cell(row=r, column=3, value=txt)
        q.cell(row=r, column=4, value=w)
        q[f'F{r}'] = f'=IF(E{r}="Oui",D{r},IF(E{r}="Partiel",D{r}/2,0))'
        q[f'G{r}'] = f'=IF(OR(E{r}="",E{r}="Sans objet"),0,D{r})'
        q[f'H{r}'] = f'=IF(OR(E{r}="",E{r}="Oui",E{r}="Sans objet"),"","{act}")'
        q.cell(row=r, column=9, value=ref)
        dv.add(f'E{r}')
        style_row(q, r, len(cols), input_cols={5}, height=32)
    last = 5 + len(QUESTIONS)
    q.conditional_formatting.add(f'E6:E{last}', CellIsRule(operator='equal', formula=['"Non"'], font=Font(color=RED, bold=True)))
    q.conditional_formatting.add(f'E6:E{last}', CellIsRule(operator='equal', formula=['"Partiel"'], font=Font(color='9A620C', bold=True)))
    q.conditional_formatting.add(f'E6:E{last}', CellIsRule(operator='equal', formula=['"Oui"'], font=Font(color='1E7346', bold=True)))

    s = wb.create_sheet('Synthèse')
    title(s, 'Synthèse du diagnostic', 'Indice de maîtrise par domaine et global (100 = toutes les pratiques en place). Le niveau d’exposition est indicatif.', 5)
    header(s, 5, ['Domaine', 'Points obtenus', 'Points possibles', 'Indice de maîtrise', 'Exposition'], [30, 14, 14, 16, 22])
    domains = []
    for d, *_ in QUESTIONS:
        if d not in domains:
            domains.append(d)
    rng = f'Questionnaire!$B$6:$B${last}'
    for k, d in enumerate(domains, start=6):
        s.cell(row=k, column=1, value=d)
        s[f'B{k}'] = f'=SUMIFS(Questionnaire!$F$6:$F${last},{rng},A{k})'
        s[f'C{k}'] = f'=SUMIFS(Questionnaire!$G$6:$G${last},{rng},A{k})'
        s[f'D{k}'] = f'=IF(C{k}=0,"",B{k}/C{k})'
        s[f'E{k}'] = f'=IF(D{k}="","Non évalué",IF(D{k}>=0.8,"Faible",IF(D{k}>=0.5,"Modérée","Élevée")))'
        s[f'D{k}'].number_format = '0%'
        style_row(s, k, 5, input_cols=set())
    t = 6 + len(domains)
    s.cell(row=t, column=1, value='Global')
    s[f'B{t}'] = f'=SUM(B6:B{t-1})'
    s[f'C{t}'] = f'=SUM(C6:C{t-1})'
    s[f'D{t}'] = f'=IF(C{t}=0,"",B{t}/C{t})'
    s[f'E{t}'] = f'=IF(D{t}="","Non évalué",IF(D{t}>=0.8,"Faible",IF(D{t}>=0.5,"Modérée","Élevée")))'
    s[f'D{t}'].number_format = '0%'
    style_row(s, t, 5, input_cols=set())
    for c in range(1, 6):
        s.cell(row=t, column=c).font = Font(name=F, bold=True, color=NAVY)
    s[f'A{t+1}'] = f'=COUNTIF(Questionnaire!$E$6:$E${last},"Non")&" réponse(s) « Non » et "&COUNTIF(Questionnaire!$E$6:$E${last},"Partiel")&" « Partiel » : voir les actions dans le questionnaire."'
    s[f'A{t+1}'].font = Font(name=F, italic=True, size=10)
    s.conditional_formatting.add(f'E6:E{t}', CellIsRule(operator='equal', formula=['"Élevée"'], font=Font(color='FFFFFF', bold=True), fill=PatternFill('solid', fgColor=RED)))
    s.conditional_formatting.add(f'E6:E{t}', CellIsRule(operator='equal', formula=['"Modérée"'], font=Font(color='9A620C', bold=True), fill=PatternFill('solid', fgColor='FCF4E3')))
    s.conditional_formatting.add(f'E6:E{t}', CellIsRule(operator='equal', formula=['"Faible"'], font=Font(color='1E7346', bold=True), fill=PatternFill('solid', fgColor='EBF6EF')))
    s.cell(row=t + 3, column=1, value='Seuils : indice ≥ 80 % = exposition faible ; 50 à 79 % = modérée ; < 50 % = élevée.').font = Font(name=F, size=9, italic=True)

    guide_sheet(wb, 'Mode d’emploi', 'Kit 03 · Auto-diagnostic de l’exposition au contrôle a posteriori', [
        ('h', 'À quoi sert ce questionnaire'),
        ('p', "Mesurer, en 32 questions, la capacité de l'entreprise à justifier ses déclarations lors d'un contrôle a posteriori. Après la mainlevée, l'administration peut réviser les déclarations, contrôler les documents et vérifier les marchandises encore présentables, dans un délai de quatre ans, qui ne court qu'à partir de la découverte en cas de fraude (CDII, art. 86 bis, 99 bis et 99 ter)."),
        ('h', 'Mode d’emploi'),
        ('p', "1. Onglet « Questionnaire » : répondez dans la colonne jaune. Chaque réponse autre que « Oui » fait apparaître l'action recommandée."),
        ('p', "2. Onglet « Synthèse » : indice de maîtrise par domaine et global, et niveau d'exposition."),
        ('p', "3. Refaites le diagnostic après chaque action corrective, et au moins une fois par an (audit interne)."),
        ('h', 'Les règles à retenir'),
        ('p', "Révélation volontaire : le déclarant qui révèle lui-même des inexactitudes dans les 30 jours suivant la mainlevée peut être dispensé de tout ou partie des pénalités, sauf s'il a déjà été informé d'un contrôle (CDII, art. 78-3°)."),
        ('p', "Conservation : registres et documents conservés 4 ans ; le défaut de conservation est puni d'une amende égale à la valeur des marchandises, le refus de communication de 30 000 à 60 000 DH (CDII, art. 42, 285-15°, 297 bis)."),
        ('p', "Ce questionnaire est un outil d'auto-évaluation : il ne remplace ni un audit externe ni l'avis d'un expert sur un dossier précis."),
        ('h', 'Sources'),
        ('p', "Guide des procédures douanières Maroc, éd. 2026, n° 2-001 à 2-061, 7-004 à 7-006, 9-002, 9-003, 9-022 ; Import Process Management, chapitre 23 et annexe D."),
    ])
    wb.active = 0
    wb.save(f'{OUT}/Kit03_Autodiagnostic_controle_a_posteriori.xlsx')


if __name__ == '__main__':
    kit01(); kit02(); kit03_registre(); kit03_diag()
    print(sorted(os.listdir(OUT)))
